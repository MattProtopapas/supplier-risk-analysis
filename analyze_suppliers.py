"""Score suppliers, assess category risk, test dual-sourcing, and write everything to one Excel file.

Reads the validated CSV produced by validate_data.py, builds a weighted supplier
scorecard, rates the risk of each purchased category, and works out whether
adding a second supplier pays for itself. Writes a brand-new workbook (no
template) with four sheets:
  - Input Data        the validated records the analysis was run on
  - Supplier Scorecard price, lead-time variability and defect scores, total score, tier and rank
  - Category Risk     incumbent vs. best alternative, expected disruption loss, risk rating, action
  - Dual-Sourcing     best single supplier vs. best two-supplier split per category, and the break-even outage probability

Scorecard (each criterion scored 0-100, then weighted; default weights 40 / 30 / 30):
  - Price score:             100 * cheapest price in the category / this supplier's price
  - Lead-time variability:   100 * (1 - min(CV / 0.5, 1)), where CV = lead-time std dev / mean lead time
  - Defect score:            100 * (1 - min(defect_rate / 8%, 1))
  - Tiers:                   A >= 75 (preferred), B >= 60 (approved), C below (watch list)

Annual cost of a sourcing set-up (demand D units/yr, daily demand d = D / 365):
  - Purchase cost:           sum over suppliers of share * D * price
  - Defect cost:             purchase share * D * defect_rate * price * defect_cost_multiplier
  - Safety stock holding:    z * d * lead-time std dev * average price * holding_rate
                             (with two suppliers the std dev is sqrt(w1^2 s1^2 + w2^2 s2^2): diversification)
  - Expected disruption loss: outage probability * d * outage_days * price * shortage_multiplier
                             (with two suppliers: both fail p1*p2 -> full loss; one fails -> only the share the
                             survivor cannot cover with surge capacity)
  - Dual-sourcing overhead:  fixed yearly cost of qualifying and managing a second supplier
Dual-sourcing pays off when its total annual cost is below the single-source total. The single-source
baseline is the lowest-total-cost supplier in the category (not necessarily the incumbent), so a saving
reflects spreading risk rather than just switching to a cheaper supplier. The tool then searches every
alternate supplier and a grid of volume splits (50-95% to the primary) for the cheapest dual set-up.
The break-even probability is the primary supplier's yearly outage probability above which dual-sourcing wins.
"""

import argparse
import csv
import math
import os
from statistics import NormalDist

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from supplier_common import CSV_FIELDS, OUTPUT_WORKBOOK_PATH, VALIDATED_DATA_PATH, Supplier

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
RISK_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
GOOD_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

SCORECARD_FIELDS = [
    "supplier_id",
    "supplier_name",
    "category",
    "country",
    "is_incumbent",
    "unit_price",
    "effective_unit_cost",
    "lead_time_mean_days",
    "lead_time_cv",
    "defect_rate_pct",
    "disruption_prob_pct",
    "price_score",
    "lead_time_score",
    "defect_score",
    "total_score",
    "tier",
    "rank_in_category",
]

CATEGORY_FIELDS = [
    "category",
    "qualified_suppliers",
    "incumbent",
    "incumbent_score",
    "incumbent_tier",
    "incumbent_rank",
    "best_supplier",
    "best_score",
    "incumbent_disruption_prob_pct",
    "annual_spend",
    "expected_disruption_loss",
    "risk_rating",
    "action",
]

DUAL_FIELDS = [
    "category",
    "primary",
    "alternate",
    "primary_share_pct",
    "single_source_cost",
    "dual_source_cost",
    "annual_saving",
    "pays_off",
    "single_disruption_loss",
    "dual_disruption_loss",
    "primary_disruption_prob_pct",
    "break_even_disruption_prob_pct",
    "recommendation",
]


class Assumptions:
    """Business assumptions behind the cost model (all overridable from the command line)."""

    def __init__(self, args: argparse.Namespace):
        self.weights = (args.weight_price, args.weight_lead_time, args.weight_defects)
        self.defect_cost_multiplier = args.defect_cost_multiplier
        self.holding_rate = args.holding_rate
        self.z = NormalDist().inv_cdf(args.service_level)
        self.outage_days = args.outage_days
        self.shortage_multiplier = args.shortage_multiplier
        self.surge = args.surge_capacity
        self.dual_fixed_cost = args.dual_fixed_cost


def read_suppliers(path: str) -> list[Supplier]:
    with open(path, newline="", encoding="utf-8") as f:
        return [Supplier(**row) for row in csv.DictReader(f)]


def tier_for(score: float) -> str:
    if score >= 75:
        return "A"
    if score >= 60:
        return "B"
    return "C"


def build_scorecard(suppliers: list[Supplier], a: Assumptions) -> list[dict]:
    min_price: dict[str, float] = {}
    for s in suppliers:
        min_price[s.category] = min(min_price.get(s.category, math.inf), s.unit_price)

    w_price, w_lt, w_def = a.weights
    total_w = (w_price + w_lt + w_def) or 1.0

    rows = []
    for s in suppliers:
        cv = s.lead_time_std_days / s.lead_time_mean_days
        price_score = 100 * min_price[s.category] / s.unit_price
        lt_score = 100 * (1 - min(cv / 0.5, 1))
        defect_score = 100 * (1 - min(s.defect_rate / 0.08, 1))
        total = (w_price * price_score + w_lt * lt_score + w_def * defect_score) / total_w
        rows.append(
            {
                "supplier_id": s.supplier_id,
                "supplier_name": s.supplier_name,
                "category": s.category,
                "country": s.country,
                "is_incumbent": s.is_incumbent,
                "unit_price": s.unit_price,
                "effective_unit_cost": round(s.unit_price * (1 + s.defect_rate * a.defect_cost_multiplier), 2),
                "lead_time_mean_days": s.lead_time_mean_days,
                "lead_time_cv": round(cv, 2),
                "defect_rate_pct": round(s.defect_rate * 100, 2),
                "disruption_prob_pct": round(s.disruption_prob * 100, 1),
                "price_score": round(price_score, 1),
                "lead_time_score": round(lt_score, 1),
                "defect_score": round(defect_score, 1),
                "total_score": round(total, 1),
                "tier": tier_for(total),
            }
        )

    rows.sort(key=lambda r: (r["category"], -r["total_score"]))
    rank = 0
    last_category = None
    for r in rows:
        rank = 1 if r["category"] != last_category else rank + 1
        r["rank_in_category"] = rank
        last_category = r["category"]
    return rows


def group_by_category(suppliers: list[Supplier]) -> dict[str, list[Supplier]]:
    groups: dict[str, list[Supplier]] = {}
    for s in suppliers:
        groups.setdefault(s.category, []).append(s)
    return groups


def pick_incumbent(group: list[Supplier]) -> Supplier:
    return next((s for s in group if s.is_incumbent), group[0])


def disruption_loss_base(s: Supplier, a: Assumptions) -> float:
    """Cost of losing the whole category supply for the full outage, before multiplying by probability."""
    daily = s.annual_volume_units / 365
    return daily * a.outage_days * s.unit_price * a.shortage_multiplier


def single_source_cost(s: Supplier, a: Assumptions, p_disrupt: float | None = None) -> tuple[float, float]:
    """Return (total annual cost, expected disruption loss) when buying everything from one supplier."""
    p = s.disruption_prob if p_disrupt is None else p_disrupt
    demand = s.annual_volume_units
    daily = demand / 365

    purchase = demand * s.unit_price
    defects = purchase * s.defect_rate * a.defect_cost_multiplier
    holding = a.z * daily * s.lead_time_std_days * s.unit_price * a.holding_rate
    disruption = p * disruption_loss_base(s, a)
    return purchase + defects + holding + disruption, disruption


def dual_source_cost(
    primary: Supplier, alt: Supplier, w: float, a: Assumptions, p_disrupt: float | None = None
) -> tuple[float, float]:
    """Return (total annual cost, expected disruption loss) when splitting volume w / (1 - w) between two suppliers."""
    p1 = primary.disruption_prob if p_disrupt is None else p_disrupt
    p2 = alt.disruption_prob
    demand = primary.annual_volume_units
    daily = demand / 365

    purchase = demand * (w * primary.unit_price + (1 - w) * alt.unit_price)
    defects = demand * (
        w * primary.unit_price * primary.defect_rate + (1 - w) * alt.unit_price * alt.defect_rate
    ) * a.defect_cost_multiplier
    sigma = math.sqrt((w * primary.lead_time_std_days) ** 2 + ((1 - w) * alt.lead_time_std_days) ** 2)
    avg_price = purchase / demand
    holding = a.z * daily * sigma * avg_price * a.holding_rate

    full_loss = daily * a.outage_days * avg_price * a.shortage_multiplier
    uncovered = 1 - a.surge  # share of a failed supplier's volume the survivor cannot absorb
    disruption = full_loss * (p1 * p2 + p1 * (1 - p2) * w * uncovered + p2 * (1 - p1) * (1 - w) * uncovered)

    return purchase + defects + holding + disruption + a.dual_fixed_cost, disruption


def best_dual_setup(group: list[Supplier], primary: Supplier, a: Assumptions):
    """Cheapest (alternate, primary share, cost, disruption loss) across all alternates and split ratios."""
    best = None
    for alt in group:
        if alt.supplier_id == primary.supplier_id:
            continue
        for step in range(10, 20):  # primary share 50% .. 95%
            w = step / 20
            cost, disruption = dual_source_cost(primary, alt, w, a)
            if best is None or cost < best[2]:
                best = (alt, w, cost, disruption)
    return best


def break_even_disruption_prob(primary: Supplier, alt: Supplier, w: float, a: Assumptions) -> float | None:
    """Primary-supplier outage probability above which the dual set-up is cheaper (the gap is linear in p)."""
    gap0 = dual_source_cost(primary, alt, w, a, 0.0)[0] - single_source_cost(primary, a, 0.0)[0]
    gap1 = dual_source_cost(primary, alt, w, a, 1.0)[0] - single_source_cost(primary, a, 1.0)[0]
    if gap0 <= 0:
        return 0.0  # dual-sourcing wins even if the primary never fails
    if gap1 >= 0:
        return None  # dual-sourcing never wins, whatever the outage probability
    return gap0 / (gap0 - gap1)


def build_dual_sourcing(groups: dict[str, list[Supplier]], a: Assumptions) -> list[dict]:
    """Compare the best single supplier against the best two-supplier split, per category.

    The single-source baseline is the lowest-total-cost supplier in the category, so any saving
    comes from spreading risk, not from simply switching to a cheaper supplier.
    """
    rows = []
    for category, group in sorted(groups.items()):
        primary = min(group, key=lambda s: single_source_cost(s, a)[0])
        single_total, single_loss = single_source_cost(primary, a)
        setup = best_dual_setup(group, primary, a)

        if setup is None:
            rows.append(
                {
                    "category": category,
                    "primary": primary.supplier_name,
                    "alternate": "none qualified",
                    "single_source_cost": round(single_total),
                    "single_disruption_loss": round(single_loss),
                    "primary_disruption_prob_pct": round(primary.disruption_prob * 100, 1),
                    "pays_off": False,
                    "recommendation": "Sole-sourced - qualify a second supplier to assess dual-sourcing",
                }
            )
            continue

        alt, w, dual_total, dual_loss = setup
        saving = single_total - dual_total
        break_even = break_even_disruption_prob(primary, alt, w, a)
        rows.append(
            {
                "category": category,
                "primary": primary.supplier_name,
                "alternate": alt.supplier_name,
                "primary_share_pct": round(w * 100),
                "single_source_cost": round(single_total),
                "dual_source_cost": round(dual_total),
                "annual_saving": round(saving),
                "pays_off": saving > 0,
                "single_disruption_loss": round(single_loss),
                "dual_disruption_loss": round(dual_loss),
                "primary_disruption_prob_pct": round(primary.disruption_prob * 100, 1),
                "break_even_disruption_prob_pct": "never" if break_even is None else round(break_even * 100, 1),
                "recommendation": (
                    f"Dual-source: {w:.0%} {primary.supplier_name} / {1 - w:.0%} {alt.supplier_name}"
                    if saving > 0
                    else "Stay single-source: the extra cost outweighs the risk reduction"
                ),
            }
        )
    return rows


def build_category_risk(scorecard: list[dict], groups: dict[str, list[Supplier]], a: Assumptions) -> list[dict]:
    by_id = {r["supplier_id"]: r for r in scorecard}
    rows = []
    for category, group in sorted(groups.items()):
        incumbent = pick_incumbent(group)
        inc_row = by_id[incumbent.supplier_id]
        best_row = max((by_id[s.supplier_id] for s in group), key=lambda r: r["total_score"])
        _, expected_loss = single_source_cost(incumbent, a)

        if len(group) == 1 or inc_row["tier"] == "C" or incumbent.disruption_prob >= 0.15:
            rating = "High"
        elif inc_row["tier"] == "B" or incumbent.disruption_prob >= 0.08:
            rating = "Medium"
        else:
            rating = "Low"

        if len(group) == 1:
            action = "Sole source: qualify a backup supplier"
        elif best_row["supplier_id"] != incumbent.supplier_id and best_row["total_score"] - inc_row["total_score"] >= 5:
            action = f"Review incumbent: {best_row['supplier_name']} scores {best_row['total_score'] - inc_row['total_score']:.0f} points higher"
        elif rating == "High":
            action = "Mitigate: dual-source or add safety stock"
        else:
            action = "Maintain"

        rows.append(
            {
                "category": category,
                "qualified_suppliers": len(group),
                "incumbent": incumbent.supplier_name,
                "incumbent_score": inc_row["total_score"],
                "incumbent_tier": inc_row["tier"],
                "incumbent_rank": inc_row["rank_in_category"],
                "best_supplier": best_row["supplier_name"],
                "best_score": best_row["total_score"],
                "incumbent_disruption_prob_pct": round(incumbent.disruption_prob * 100, 1),
                "annual_spend": round(incumbent.annual_volume_units * incumbent.unit_price),
                "expected_disruption_loss": round(expected_loss),
                "risk_rating": rating,
                "action": action,
            }
        )
    return rows


def autosize_columns(ws) -> None:
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(length + 2, 70)


def write_sheet(ws, fieldnames: list[str], rows: list[dict]) -> None:
    ws.append(fieldnames)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append([row.get(f, "") for f in fieldnames])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(fieldnames))}{len(rows) + 1}"
    autosize_columns(ws)


def highlight_rows(ws, fieldnames: list[str], column: str, predicate, fill) -> None:
    col = fieldnames.index(column)
    for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if predicate(row_cells[col].value):
            for cell in row_cells:
                cell.fill = fill


def build_workbook(input_rows: list[dict], scorecard: list[dict], category_risk: list[dict], dual: list[dict]) -> Workbook:
    wb = Workbook()

    ws_input = wb.active
    ws_input.title = "Input Data"
    write_sheet(ws_input, CSV_FIELDS, input_rows)

    ws_score = wb.create_sheet("Supplier Scorecard")
    write_sheet(ws_score, SCORECARD_FIELDS, scorecard)
    highlight_rows(ws_score, SCORECARD_FIELDS, "tier", lambda v: v == "C", RISK_FILL)

    ws_risk = wb.create_sheet("Category Risk")
    write_sheet(ws_risk, CATEGORY_FIELDS, category_risk)
    highlight_rows(ws_risk, CATEGORY_FIELDS, "risk_rating", lambda v: v == "High", RISK_FILL)

    ws_dual = wb.create_sheet("Dual-Sourcing")
    write_sheet(ws_dual, DUAL_FIELDS, dual)
    highlight_rows(ws_dual, DUAL_FIELDS, "pays_off", lambda v: v is True, GOOD_FILL)

    return wb


def main() -> None:
    parser = argparse.ArgumentParser(description="Score suppliers, rate category risk and test dual-sourcing; write results to Excel.")
    parser.add_argument("--input", type=str, default=VALIDATED_DATA_PATH, help="Path to the validated CSV")
    parser.add_argument("--output", type=str, default=OUTPUT_WORKBOOK_PATH, help="Path to write the Excel workbook to")
    parser.add_argument("--weight-price", type=float, default=0.40, help="Scorecard weight for price (default: 0.40)")
    parser.add_argument("--weight-lead-time", type=float, default=0.30, help="Scorecard weight for lead-time variability (default: 0.30)")
    parser.add_argument("--weight-defects", type=float, default=0.30, help="Scorecard weight for defect rate (default: 0.30)")
    parser.add_argument("--defect-cost-multiplier", type=float, default=2.0, help="Cost of a defective unit as a multiple of its price (default: 2.0)")
    parser.add_argument("--holding-rate", type=float, default=0.25, help="Yearly inventory holding cost as a share of unit value (default: 0.25)")
    parser.add_argument("--service-level", type=float, default=0.95, help="Service level used to size safety stock (default: 0.95)")
    parser.add_argument("--outage-days", type=float, default=30, help="Length of a major supplier outage in days (default: 30)")
    parser.add_argument("--shortage-multiplier", type=float, default=1.5, help="Cost of a unit not delivered, as a multiple of its price (default: 1.5)")
    parser.add_argument("--surge-capacity", type=float, default=0.5, help="Share of a failed supplier's volume the other can absorb (default: 0.5)")
    parser.add_argument("--dual-fixed-cost", type=float, default=15000, help="Yearly cost of qualifying and managing a second supplier (default: 15000)")
    args = parser.parse_args()

    a = Assumptions(args)
    suppliers = read_suppliers(args.input)
    with open(args.input, newline="", encoding="utf-8") as f:
        input_rows = list(csv.DictReader(f))

    groups = group_by_category(suppliers)
    scorecard = build_scorecard(suppliers, a)
    category_risk = build_category_risk(scorecard, groups, a)
    dual = build_dual_sourcing(groups, a)

    wb = build_workbook(input_rows, scorecard, category_risk, dual)
    os.makedirs(os.path.dirname(args.output) or ".", exist_ok=True)
    wb.save(args.output)

    high = sum(1 for r in category_risk if r["risk_rating"] == "High")
    pays = sum(1 for r in dual if r["pays_off"])
    print(f"Scored {len(suppliers)} suppliers in {len(groups)} categories -> {args.output}")
    print(f"{high} categor(ies) rated high risk; dual-sourcing pays off in {pays} of {len(dual)}.")


if __name__ == "__main__":
    main()
